#!/usr/bin/env python3
"""Build the PXD037708 protein-intensity vignette fixture from PRIDE."""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import io
import math
import tempfile
import urllib.request
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as error:  # pragma: no cover - regeneration dependency
    raise SystemExit(
        "Install the regeneration dependency with `python -m pip install openpyxl`."
    ) from error


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
SOURCE_URL = (
    "https://ftp.pride.ebi.ac.uk/pride/data/archive/2025/05/PXD037708/"
    "20221024_TAU_EV64_PRIDE.xlsx"
)
SOURCE_SHA256 = "a018b3a20904f07f49928014056d0c005dab14b933e9d4b6cceadd28764cc34b"
DEFAULT_MATRIX_OUTPUT = (
    REPOSITORY_ROOT / "inst" / "extdata" / "PXD037708_brain_EV_log2.tsv.gz"
)
DEFAULT_SAMPLE_OUTPUT = (
    REPOSITORY_ROOT / "inst" / "extdata" / "PXD037708_brain_EV_samples.csv"
)
REMOVED_SAMPLES = {"A_F8", "E_F8", "F_F8"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def acquire_workbook(path: Path | None, temporary_directory: Path) -> Path:
    workbook = path.resolve() if path else temporary_directory / Path(SOURCE_URL).name
    if path is None:
        print(f"Downloading {SOURCE_URL}")
        urllib.request.urlretrieve(SOURCE_URL, workbook)
    observed = sha256(workbook)
    if observed != SOURCE_SHA256:
        raise SystemExit(
            f"PXD037708 workbook checksum mismatch: expected {SOURCE_SHA256}, got {observed}."
        )
    return workbook


def sample_order(sample: str) -> tuple[int, str]:
    donor, fraction = sample.split("_")
    return int(fraction.removeprefix("F")), donor


def log2_value(value: object) -> str:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return ""
    if not math.isfinite(numeric) or numeric <= 0:
        return ""
    return f"{math.log2(numeric):.6f}"


def write_matrix(workbook_path: Path, output: Path) -> tuple[list[str], int]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook["PROTEINS"]
    rows = worksheet.iter_rows(values_only=True)
    header = list(next(rows))
    next(rows)  # Patient labels beneath the sample header.

    source_samples = list(header[4:])
    retained_indices = [
        index
        for index, sample in enumerate(source_samples)
        if sample not in REMOVED_SAMPLES
    ]
    retained_indices.sort(key=lambda index: sample_order(source_samples[index]))
    samples = [source_samples[index] for index in retained_indices]

    output.parent.mkdir(parents=True, exist_ok=True)
    row_count = 0
    with output.open("wb") as raw_handle:
        with gzip.GzipFile(
            filename="",
            mode="wb",
            fileobj=raw_handle,
            mtime=0,
        ) as compressed_handle:
            with io.TextIOWrapper(
                compressed_handle,
                encoding="utf-8",
                newline="",
            ) as text_handle:
                writer = csv.writer(text_handle, delimiter="\t", lineterminator="\n")
                writer.writerow(
                    [
                        "Protein.Group",
                        "Protein.Names",
                        "Genes",
                        "First.Protein.Description",
                        *samples,
                    ]
                )
                for row in rows:
                    metadata = [row[index] if row[index] is not None else "" for index in range(4)]
                    intensities = [log2_value(row[4 + index]) for index in retained_indices]
                    writer.writerow([*metadata, *intensities])
                    row_count += 1

    workbook.close()
    if row_count != 6107 or len(samples) != 61:
        raise SystemExit(
            f"Expected 6107 proteins and 61 samples, wrote {row_count} and {len(samples)}."
        )
    return samples, row_count


def write_samples(samples: list[str], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["sample", "donor", "fraction", "condition"])
        for sample in samples:
            donor, fraction_label = sample.split("_")
            fraction = int(fraction_label.removeprefix("F"))
            writer.writerow([sample, donor, fraction, f"Fraction {fraction}"])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--workbook",
        type=Path,
        help="Existing checksum-matched PRIDE workbook; download when omitted.",
    )
    parser.add_argument("--matrix-output", type=Path, default=DEFAULT_MATRIX_OUTPUT)
    parser.add_argument("--sample-output", type=Path, default=DEFAULT_SAMPLE_OUTPUT)
    arguments = parser.parse_args()

    with tempfile.TemporaryDirectory(prefix="imputefinder-pxd037708-") as temporary:
        workbook = acquire_workbook(arguments.workbook, Path(temporary))
        samples, row_count = write_matrix(workbook, arguments.matrix_output.resolve())
        write_samples(samples, arguments.sample_output.resolve())

    print(
        f"Wrote {row_count} proteins × {len(samples)} samples to "
        f"{arguments.matrix_output.resolve()}."
    )
    print(f"Wrote sample metadata to {arguments.sample_output.resolve()}.")


if __name__ == "__main__":
    main()
